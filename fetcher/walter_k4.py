"""WALTER K4 + avgiftsrapport: CSV och PDF ur state.json (journalen är facit).

  python fetcher/walter_k4.py

Skriver till reports/ (genereras, checkas inte in):
  K4_<år>_WALTER_PAPER.csv              K4 sektion D, en rad per försäljning
  K4-WALTER-<år>-PAPER.xlsx             samma rader som riktig Excel: datumceller, breda kolumner, summor
  K4_<år>_WALTER_PAPER.pdf              samma sak, snyggt
  Avgiftsrapport_<år>_WALTER_PAPER.pdf  avgiftsmodellen, betalda avgifter, uttagsexempel, portföljens väg

År = senaste försäljningens år (annars innevarande år) – samma regel som knapparna i app.js.
Allt är PAPER: vattenstämpel "PAPER – simulering, inte för Skatteverket eller bank".
När WALTER går LIVE ska raderna i stället läsas från Krakens ledger.csv och stämpeln tas bort.
"""

import csv
import json
import sys
from datetime import date, datetime
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parent.parent
STATE = ROOT / "data" / "state.json"
CACHE = ROOT / "data" / "cache.json"
OUT = ROOT / "reports"
WATERMARK = "PAPER – simulering, inte för Skatteverket eller bank"
GAIN_TAX, LOSS_DEDUCT = 0.30, 0.70  # vinst beskattas 30 %, förlust dras av till 70 % (K4 sektion D)

INK, MUTED, ACCENT, RED = colors.HexColor("#080A12"), colors.HexColor("#6B6F80"), colors.HexColor("#5741F8"), colors.HexColor("#C8102E")


# ---------- Format (svenska, som i app.js – men "-" i stället för "−", PDF-fonten saknar det tecknet) ----------
def num(v, d=0):
    s = f"{abs(v):,.{d}f}".replace(",", " ").replace(".", ",")
    return ("-" if v < 0 else "") + s


def kr(v, d=2):
    return num(v, d) + " kr"


def read_json(path, default):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default


# ---------- K4 ----------
def k4_rows(state, year):
    """Sektion D: försäljningspris efter avgifter, omkostnadsbelopp, vinst eller förlust. Hela kronor på blanketten."""
    rows = []
    for r in state.get("k4", []):
        if not str(r.get("saljdatum", "")).startswith(str(year)):
            continue
        fee = (r.get("trading_fee_sek") or 0) + (r.get("withdrawal_fee_sek") or 0)
        sale = r["forsaljningspris_sek"] - fee
        result = sale - r["omkostnad_sek"]
        rows.append({**r, "fee": fee, "sale_net": sale, "result": result,
                     "sale_k4": round(sale), "cost_k4": round(r["omkostnad_sek"])})
    return rows


def k4_summary(rows):
    gains = sum(r["result"] for r in rows if r["result"] > 0)
    losses = -sum(r["result"] for r in rows if r["result"] < 0)
    deductible = losses * LOSS_DEDUCT
    return {"gains": gains, "losses": losses, "deductible": deductible,
            "tax": max(0.0, (gains - deductible) * GAIN_TAX)}


def write_csv(rows, path):
    with path.open("w", newline="", encoding="utf-8-sig") as f:  # BOM + semikolon = öppnas rätt i svensk Excel
        w = csv.writer(f, delimiter=";")
        # Datum och exakta belopp först. Vinst/Förlust = försäljningspris − avgift − omkostnad.
        # Sist: hela kronor som på blanketten, och PAPER-stämpeln.
        w.writerow(["Köpdatum", "Säljdatum", "Antal/Belopp", "Försäljningspris (kr)", "Omkostnadsbelopp (kr)",
                    "Vinst/Förlust (kr)", "Avgift (kr)", "Beteckning", "Typ",
                    "K4 försäljningspris (hela kr)", "K4 omkostnad (hela kr)", "Status"])
        for r in rows:
            w.writerow([r["kopdatum"], r["saljdatum"], f"{num(r['antal'], 8)} {r['beteckning'].split('(')[-1].rstrip(')')}",
                        num(r["forsaljningspris_sek"], 2), num(r["omkostnad_sek"], 2), num(r["result"], 2),
                        num(r["fee"], 2), r["beteckning"], r.get("typ", ""),
                        r["sale_k4"], r["cost_k4"], WATERMARK])


def write_xlsx(rows, summ, year, path):
    """Riktig Excel: datum som datumceller (ÅÅÅÅ-MM-DD), belopp som tal, kolumner breda nog för att slippa ######."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = f"K4 {year}"
    # (rubrik, bredd i tecken, talformat)
    cols = [("Köpdatum", 12, "yyyy-mm-dd"), ("Säljdatum", 12, "yyyy-mm-dd"), ("Antal", 15, "0.00000000"),
            ("Försäljningspris (kr)", 12, "#,##0.00"), ("Omkostnadsbelopp (kr)", 12, "#,##0.00"),
            ("Vinst/Förlust (kr)", 12, "#,##0.00"), ("Avgift (kr)", 12, "#,##0.00"), ("Beteckning", 16, None),
            ("Typ", 16, None), ("K4 försäljningspris (hela kr)", 12, "#,##0"), ("K4 omkostnad (hela kr)", 12, "#,##0")]
    ws.append([WATERMARK.upper()])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws["A1"].font = Font(bold=True, color="C8102E")
    ws.append([c[0] for c in cols])
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5741F8")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([date.fromisoformat(r["kopdatum"]), date.fromisoformat(r["saljdatum"]), r["antal"],
                   r["forsaljningspris_sek"], r["omkostnad_sek"], round(r["result"], 2), round(r["fee"], 2),
                   r["beteckning"], r.get("typ", ""), r["sale_k4"], r["cost_k4"]])
    for i, (_, width, fmt) in enumerate(cols, start=1):
        letter = ws.cell(row=2, column=i).column_letter
        ws.column_dimensions[letter].width = width + 2  # lite luft, så att datum och belopp aldrig blir ######
        if fmt:
            for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=i, max_col=i):
                row[0].number_format = fmt
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "A3"

    ws.append([])
    for label, value in [("Summa vinster (kr)", summ["gains"]), ("Summa förluster (kr)", summ["losses"]),
                         (f"Avdragsgill förlust {int(LOSS_DEDUCT * 100)} % (kr)", summ["deductible"]),
                         (f"Uppskattad skatt {int(GAIN_TAX * 100)} % (kr)", summ["tax"])]:
        ws.append([label, None, None, None, None, round(value, 2)])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=6).number_format = "#,##0.00"
    wb.save(path)


# ---------- PDF ----------
def stamp(canvas, doc):
    """Vattenstämpel på varje sida: diagonalt över sidan + röd rad överst och sidfot."""
    w, h = A4
    canvas.saveState()
    canvas.setFillColor(colors.Color(0.78, 0.06, 0.18, alpha=0.16))
    canvas.setFont("Helvetica-Bold", 44)
    canvas.translate(w / 2, h / 2)
    canvas.rotate(35)
    canvas.drawCentredString(0, 22, "PAPER – simulering")
    canvas.setFont("Helvetica-Bold", 22)
    canvas.drawCentredString(0, -14, "inte för Skatteverket eller bank")
    canvas.restoreState()
    canvas.saveState()
    canvas.setFillColor(RED)
    canvas.setFont("Helvetica-Bold", 9)
    canvas.drawCentredString(w / 2, h - 12 * mm, WATERMARK.upper())
    canvas.setFillColor(MUTED)
    canvas.setFont("Helvetica", 8)
    canvas.drawString(18 * mm, 10 * mm, f"WALTER · skapad {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    canvas.drawRightString(w - 18 * mm, 10 * mm, f"Sida {doc.page}")
    canvas.restoreState()


STYLES = getSampleStyleSheet()
H1 = ParagraphStyle("h1", parent=STYLES["Title"], fontName="Helvetica-Bold", fontSize=20, textColor=INK, alignment=0, spaceAfter=4)
H2 = ParagraphStyle("h2", parent=STYLES["Heading2"], fontName="Helvetica-Bold", fontSize=12.5, textColor=ACCENT, spaceBefore=12, spaceAfter=6)
BODY = ParagraphStyle("body", parent=STYLES["BodyText"], fontName="Helvetica", fontSize=10, leading=14, textColor=INK)
SMALL = ParagraphStyle("small", parent=BODY, fontSize=8.5, leading=11, textColor=MUTED)


def table(data, widths, right_cols=(), total=False):
    t = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    style = [
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 9),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F2FE")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDDCF0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for c in right_cols:
        style.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    if total:
        style += [("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 9), ("LINEABOVE", (0, -1), (-1, -1), 0.8, INK)]
    t.setStyle(TableStyle(style))
    return t


class StampedDoc(SimpleDocTemplate):
    def afterPage(self):
        stamp(self.canv, self)  # efter innehållet, så att tabellerna inte täcker stämpeln


def build_pdf(path, story):
    StampedDoc(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
               topMargin=20 * mm, bottomMargin=18 * mm, title=path.stem, author="WALTER").build(story)


def k4_pdf(rows, summ, year, path):
    story = [Paragraph(f"K4 {year} · sektion D", H1),
             Paragraph("Kryptovalutor (övriga tillgångar). Underlag från WALTER:s paper-journal i state.json.", SMALL),
             Paragraph("Försäljningar", H2)]
    if rows:
        data = [["Antal", "Beteckning", "Köpt", "Såld", "Försäljningspris*", "Omkostnad", "Vinst", "Förlust"]]
        for r in rows:
            data.append([num(r["antal"], 8), r["beteckning"], r["kopdatum"], r["saljdatum"],
                         num(r["sale_k4"]), num(r["cost_k4"]),
                         num(max(0, r["sale_k4"] - r["cost_k4"])), num(max(0, r["cost_k4"] - r["sale_k4"]))])
        story.append(table(data, [24, 26, 20, 20, 26, 20, 19, 19], right_cols=(0, 4, 5, 6, 7)))
        story.append(Paragraph("* Efter avdrag för försäljningsavgifter. Hela kronor, som på blanketten.", SMALL))
        story.append(Paragraph("Exakt uträkning", H2))
        data = [["Köpt", "Såld", "Säljpris", "Avgift", "Netto", "Omkostnad", "Vinst/Förlust"]]
        for r in rows:
            data.append([r["kopdatum"], r["saljdatum"], kr(r["forsaljningspris_sek"]), kr(-r["fee"]),
                         kr(r["sale_net"]), kr(-r["omkostnad_sek"]), kr(r["result"])])
        story.append(table(data, [22, 22, 23, 21, 24, 25, 27], right_cols=(2, 3, 4, 5, 6)))
        notes = [r["not"] for r in rows if r.get("not")]
        if notes:
            story.append(Spacer(1, 4))
            story += [Paragraph(n, SMALL) for n in notes]
    else:
        story.append(Paragraph(f"Inga försäljningar {year}.", BODY))

    story.append(Paragraph("Sammanställning", H2))
    data = [["", "Belopp"],
            ["Summa vinster", kr(summ["gains"])],
            ["Summa förluster", kr(summ["losses"])],
            [f"Avdragsgill förlust ({int(LOSS_DEDUCT * 100)} %)", kr(summ["deductible"])],
            [f"Uppskattad skatt ({int(GAIN_TAX * 100)} % av vinst efter avdrag)", kr(summ["tax"])]]
    story.append(table(data, [120, 54], right_cols=(1,), total=True))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Vinster och förluster står på separata rader. Omkostnadsbeloppet räknas med genomsnittsmetoden "
                           "(Skatteverkets regel för kryptovalutor) inklusive köpavgift. Bankens valutaväxling efter försäljningen hör inte till K4.", SMALL))
    build_pdf(path, story)


def withdrawal_example(fees, eursek, netto=50000):
    """Samma formel som Mitt mål i app.js: sälj (taker + slippage) → EUR → SEPA → banken växlar (FX)."""
    t, s, f = fees["takerPct"] / 100, fees["slippagePct"] / 100, fees["fxPct"] / 100
    sepa = fees["sepaEur"] * eursek
    before_fx = netto / (1 - f)
    brutto = (before_fx + sepa) / ((1 - t) * (1 - s))
    return brutto, [("Kraken taker", brutto * t), ("Slippage", brutto * (1 - t) * s), ("SEPA-uttag", sepa),
                    ("Valutaväxling", before_fx - netto)]


def fees_pdf(state, cache, rows, year, path):
    fees = state["rules"].get("withdrawalFees", {})
    pf = state["portfolio"]
    story = [Paragraph(f"Avgiftsrapport {year}", H1),
             Paragraph("Vilka avgifter WALTER räknar med, vad som har betalats och hur pengarna rör sig. "
                       "Underlag: state.json och senast hämtade kurser.", SMALL),
             Paragraph("Avgiftsmodell", H2)]
    data = [["Steg", "Avgift", "När den tas"],
            ["Kraken taker", f"{num(fees.get('takerPct', 0), 2)} %", "Vid försäljning på Kraken"],
            ["Slippage", f"{num(fees.get('slippagePct', 0), 1)} %", "Skillnad mot visat pris vid försäljning"],
            ["SEPA-uttag", f"€{num(fees.get('sepaEur', 0), 2)}", "Per uttag från Kraken till banken"],
            ["Valutaväxling", f"{num(fees.get('fxPct', 0), 1)} %", "När banken växlar EUR till SEK"]]
    story.append(table(data, [40, 26, 108], right_cols=(1,)))

    story.append(Paragraph(f"Betalda avgifter {year}", H2))
    if rows:
        data = [["Datum", "Händelse", "Belopp", "Avgift"]]
        for r in rows:
            data.append([r["saljdatum"], f"Sålt {r['beteckning']} – {r.get('typ', '')}", kr(r["forsaljningspris_sek"]), kr(r["fee"])])
        data.append(["", "Summa", "", kr(sum(r["fee"] for r in rows))])
        story.append(table(data, [26, 92, 28, 28], right_cols=(2, 3), total=True))
        story += [Paragraph(r["not"], SMALL) for r in rows if r.get("not")]
    else:
        story.append(Paragraph("Inga avgifter betalda.", BODY))

    story.append(Paragraph("Exempel: uttag på 50 000 kr netto", H2))
    eur = (cache.get("market", {}).get("eursek") or state["market"].get("eursek") or {})
    if eur.get("value") and all(k in fees for k in ("takerPct", "slippagePct", "sepaEur", "fxPct")):
        brutto, parts = withdrawal_example(fees, eur["value"])
        data = [["", "Belopp"]] + [[n, kr(v)] for n, v in parts] + \
               [["Avgifter totalt", kr(brutto - 50000)], ["Brutto som måste säljas", kr(brutto)]]
        story.append(table(data, [120, 54], right_cols=(1,), total=True))
        story.append(Paragraph(f"EURSEK {num(eur['value'], 4)} ({eur.get('asOf', 'okänt datum')}, {eur.get('source', '')}).", SMALL))
    else:
        story.append(Paragraph("EURSEK eller avgifter saknas – exemplet räknas när kursen har hämtats.", BODY))

    story.append(Paragraph("Portföljens väg", H2))
    cc = pf.get("capitalChange")
    data = [["Datum", "Värde", "Varav kontant", "Not"]]
    for h in pf.get("history", []):
        if cc and h["date"] >= cc["date"] and not any(r[3] == "insättning" for r in data):
            data.append([cc["date"], kr(cc["value"], 0), "", "insättning"])  # brottet mellan fas 1 och fas 2
        data.append([h["date"], kr(h["value"], 0), kr(h["cash"], 0) if h.get("cash") is not None else "–", h.get("note", "")])
    story.append(table(data, [32, 40, 40, 30], right_cols=(1, 2)))
    if cc:
        story.append(Paragraph(f"{cc['text']}. Före {cc['date']}: paper-test. Insättningen är ingen vinst.", SMALL))
    top_up = (cc or {}).get("manualTopUp", 0)
    deposits = pf["start"] + top_up
    gain = pf["value"] - deposits  # marknadsvinst = det som är över insättningarna
    story.append(Spacer(1, 6))
    story.append(Paragraph(f"Insättningar {kr(deposits, 0)} (startkapital {kr(pf['start'], 0)}"
                           f"{f' + manuell insättning {kr(top_up, 0)}' if top_up else ''}), nu {kr(pf['value'], 0)}. "
                           f"Marknadsvinst {'+' if gain >= 0 else ''}{kr(gain, 0)}. Paper trading – inga riktiga pengar har rört sig.", BODY))
    build_pdf(path, story)


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    state = read_json(STATE, None)
    if state is None:
        sys.exit(f"Hittar inte {STATE}")
    cache = read_json(CACHE, {})
    sales = [r.get("saljdatum", "") for r in state.get("k4", []) if r.get("saljdatum")]
    year = int(max(sales)[:4]) if sales else date.today().year
    rows = k4_rows(state, year)
    summ = k4_summary(rows)

    OUT.mkdir(exist_ok=True)
    base = f"K4_{year}_WALTER_PAPER"
    write_csv(rows, OUT / f"{base}.csv")
    write_xlsx(rows, summ, year, OUT / f"K4-WALTER-{year}-PAPER.xlsx")
    k4_pdf(rows, summ, year, OUT / f"{base}.pdf")
    fees_pdf(state, cache, rows, year, OUT / f"Avgiftsrapport_{year}_WALTER_PAPER.pdf")
    print(f"K4 {year}: {len(rows)} försäljning(ar), vinst {kr(summ['gains'])}, förlust {kr(summ['losses'])}, "
          f"uppskattad skatt {kr(summ['tax'])}. Skrev reports/.")


if __name__ == "__main__":
    main()
